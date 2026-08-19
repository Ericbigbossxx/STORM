from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from storm.adapters.cm_workbook.adapter import CmWorkbookAdapter, WorkbookContractError
from storm.structured_metrics.schema import ExceptionSeverity, Platform


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "data" / "inbox" / "US E-commerce Actual CM for 2026_vs. BP.xlsx"
CONFIG = ROOT / "config" / "structured_metrics.yaml"


@pytest.fixture(scope="module")
def adapter() -> CmWorkbookAdapter:
    return CmWorkbookAdapter(CONFIG)


@pytest.fixture(scope="module")
def official_plan(adapter: CmWorkbookAdapter):
    if not WORKBOOK.exists():
        pytest.skip("official local workbook is intentionally not committed")
    return adapter.parse(WORKBOOK, snapshot_date="2026-08-11", data_through_date="2026-08-10")


def test_explicit_platform_mapping_and_invalid_platform(adapter):
    assert adapter.map_platform("Walmart Seller") is Platform.WALMART_MP
    assert adapter.map_platform("The Home Depot Inc") is Platform.THD
    assert adapter.map_platform("Lowe's") is Platform.LOWES
    with pytest.raises(WorkbookContractError, match="invalid platform mapping"):
        adapter.map_platform("Amazon")


def test_verified_workbook_structure_and_bp_grain(official_plan):
    assert official_plan.file_hash == "89B8BF3EAD9D940833582CA01FC0065C4099E335C30EC9FEAC0D5CB2CBBC1168"
    assert official_plan.source_period == "2026-08|MTD"
    assert len(official_plan.dim_skus) == 93
    assert len(official_plan.bp_targets) == 1110
    assert len(official_plan.cm_snapshots) == 15
    assert len({(row.year, row.month, row.platform, row.canonical_sku) for row in official_plan.bp_targets}) == 1110


def test_bp_cm_includes_ddp_cogs_and_all_verified_costs(official_plan):
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook["KPI Rawdata"]
        headers = [cell.value for cell in sheet[1]]
        index = {name: position for position, name in enumerate(headers)}
        target = next(
            row for row in official_plan.bp_targets
            if row.platform is Platform.THD and row.year == 2026 and row.month == 8
        )
        source_rows = [tuple(row) for row in sheet.iter_rows(min_row=2, values_only=True)
                       if row[index["Customer"]] == "The Home Depot Inc"
                       and str(row[index["SKU"]]).strip() == target.canonical_sku
                       and row[index["Year"]] == 2026 and row[index["Month-INT"]] == 8]
        total = sum(
            row[index["TTL amount"]]
            - row[index["DDP/ALL"]]
            - row[index["Fixed cost/ALL"]]
            - row[index["MKT-Insite/ALL"]]
            - row[index["MKT-Offsite(种草)/ALL"]]
            - row[index["MKT-Offsite(Channel MKT)/ALL"]]
            - row[index["Return+Warranty/ALL"]]
            - row[index["Funding/ALL"]]
            for row in source_rows
        )
    finally:
        workbook.close()
    assert target.bp_cm == pytest.approx(total)


def test_duplicate_normalized_mapping_is_preserved_as_warning(official_plan):
    duplicate = [row for row in official_plan.exceptions if row.exception_type == "DUPLICATE_SKU_MAPPING"]
    assert len(duplicate) == 1
    assert duplicate[0].raw_key == "SK-L-WIRE"
    assert duplicate[0].severity is ExceptionSeverity.WARNING


def test_thd_primary_snapshot_excludes_dfc(official_plan):
    thd = [row for row in official_plan.cm_snapshots if row.platform is Platform.THD]
    assert len(thd) == 5
    assert all("DFC" not in row.source_range for row in thd)
    assert all(row.cm_basis == "PRIMARY_SELL_IN_EXCLUDES_DFC" for row in thd)
    assert sum(row.actual_cm for row in thd if row.actual_cm is not None) == pytest.approx(15104.261940370336)
    # The source's combined THD summary includes a -612.5 DFC CM; primary THD deliberately does not.
    assert sum(row.actual_cm for row in thd if row.actual_cm is not None) != pytest.approx(14491.761940370336)


def test_null_semantics_and_reconciliation_gate(official_plan):
    assert any(row.bp_cm_pct is None for row in official_plan.bp_targets)
    assert official_plan.reconciliations
    platform = [row for row in official_plan.reconciliations if row.level == "PLATFORM"]
    assert len(platform) == 18
    assert max(abs(row.variance or 0.0) for row in platform) <= 0.01
    bp_detail = [row for row in official_plan.reconciliations if row.level == "BP_DERIVED_VIEW_BRAND_POWER"]
    bp_platform = [row for row in official_plan.reconciliations if row.level == "BP_DERIVED_VIEW_PLATFORM"]
    assert len(bp_detail) == 45
    assert len(bp_platform) == 9
    failures = [row for row in (*bp_detail, *bp_platform) if row.result == "FAIL"]
    assert {(row.level, row.key, row.metric) for row in failures} == {
        ("BP_DERIVED_VIEW_BRAND_POWER", "WALMART_MP|Sunseeker|Robot", "bp_units"),
        ("BP_DERIVED_VIEW_BRAND_POWER", "WALMART_MP|Sunseeker|Robot", "bp_sales"),
        ("BP_DERIVED_VIEW_BRAND_POWER", "WALMART_MP|Sunseeker|Robot", "bp_cm"),
        ("BP_DERIVED_VIEW_PLATFORM", "WALMART_MP", "bp_units"),
        ("BP_DERIVED_VIEW_PLATFORM", "WALMART_MP", "bp_sales"),
        ("BP_DERIVED_VIEW_PLATFORM", "WALMART_MP", "bp_cm"),
    }
    assert all(not row.is_blocking and row.classification == "BROKEN_DERIVED_FORMULA" for row in failures)
    warnings = [row for row in official_plan.exceptions if row.exception_type == "BROKEN_DERIVED_FORMULA"]
    assert len(warnings) == 6
    assert all(row.severity is ExceptionSeverity.WARNING for row in warnings)
    assert official_plan.source_totals["coverage"] == {
        "MTD": "ACCEPTED",
        "YTD": "SOURCE_NOT_AVAILABLE",
    }


def test_data_through_date_must_match_source(adapter):
    if not WORKBOOK.exists():
        pytest.skip("official local workbook is intentionally not committed")
    with pytest.raises(WorkbookContractError, match="data_through_date mismatch"):
        adapter.parse(WORKBOOK, snapshot_date="2026-08-11", data_through_date="2026-08-09")

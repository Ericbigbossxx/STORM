"""Read source-reported CM and operating performance from ``over view``.

The reader treats coordinates only as discovered locations.  Every value is
accepted only after its section label, scenario header, dimension header, and
metric label have been validated.  CM rates and percentage-point gaps retain
the workbook's fractional numeric representation so source equality remains
auditable (for example, ``0.0345`` displays as ``3.45 pp`` for a gap).
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from .normalization import as_number, normalize_text


PHASE3_CONTRACT_VERSION = "3.0.0-business-performance"
PHASE4_CONTRACT_VERSION = "4.0.0-operating-efficiency"
METRIC_DOMAIN = "CM_BUSINESS_PERFORMANCE"
METRIC_NAME = "contribution_margin_rate"
OPERATING_DOMAIN = "OPERATING_EFFICIENCY"
OPERATING_SOURCE_LABELS = {
    "gmv": "GMV",
    "market_insight": "MKT-Insite",
    "fixed_cost": "Fixed Cost",
    "return_warranty_cost": "Return+Warranty",
    "funding": "Funding",
}
CHANNEL_MAP = {
    "HomeDepot*": ("THD", "DS"),
    "HomeDepot": ("THD", "DS"),
    "Lowes": ("Lowe's", "DS"),
    "Walmart MP": ("Walmart", "MP"),
    "Walmart DSV": ("Walmart", "DSV"),
}


@dataclass(slots=True)
class BusinessPerformanceData:
    snapshot_id: str
    source: dict[str, Any]
    period: dict[str, Any]
    records: list[dict[str, Any]]
    coverage: dict[str, Any]
    quality_issues: list[dict[str, Any]]
    reconciliation: dict[str, Any]
    operating_records: list[dict[str, Any]]
    operating_coverage: dict[str, Any]
    operating_reconciliation: dict[str, Any]
    operating_source_audit: dict[str, Any]


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest().upper()


def _text(value: Any) -> str | None:
    return normalize_text(value)


def _matches(value: Any, expected: str) -> bool:
    actual = _text(value)
    return actual is not None and actual.casefold() == expected.casefold()


def _find_cells(ws, label: str, *, min_row: int = 1, max_row: int | None = None) -> list[Any]:
    limit = max_row or ws.max_row
    return [
        cell
        for row in ws.iter_rows(min_row=min_row, max_row=limit)
        for cell in row
        if _matches(cell.value, label)
    ]


def _find_one(ws, label: str, *, min_row: int = 1, max_row: int | None = None) -> Any:
    found = _find_cells(ws, label, min_row=min_row, max_row=max_row)
    if len(found) != 1:
        refs = ", ".join(cell.coordinate for cell in found) or "none"
        raise ValueError(f"Expected one {label!r} label, found {len(found)} at {refs}")
    return found[0]


def _find_in_row(ws, row: int, label: str, *, min_col: int = 1, max_col: int | None = None) -> list[Any]:
    limit = max_col or ws.max_column
    return [
        ws.cell(row, column)
        for column in range(min_col, limit + 1)
        if _matches(ws.cell(row, column).value, label)
    ]


def _metric_row(ws, label_column: int, start_row: int, end_row: int, metric: str = "CM%") -> int:
    found = [
        row
        for row in range(start_row, end_row + 1)
        if _matches(ws.cell(row, label_column).value, metric)
    ]
    if len(found) != 1:
        raise ValueError(
            f"Expected one {metric!r} label in column {label_column}, rows {start_row}:{end_row}; "
            f"found {found}"
        )
    return found[0]


def _base_dimensions(level: str, canonical: Any) -> dict[str, Any]:
    dimensions = {
        "platform": None,
        "channel": None,
        "subchannel": None,
        "brand": None,
        "power_source": None,
        "sku": None,
    }
    if level == "CHANNEL":
        dimensions["platform"], dimensions["channel"] = canonical
    elif level == "BRAND":
        dimensions["brand"] = canonical
    elif level == "POWER_SOURCE":
        dimensions["power_source"] = canonical
    else:
        raise ValueError(f"Unsupported analysis level: {level}")
    return dimensions


def _source_sign(value: float | None) -> str:
    if value is None:
        return "MISSING"
    if value > 0:
        return "POSITIVE"
    if value < 0:
        return "NEGATIVE"
    return "ZERO"


def _cost_sign_convention(values: Iterable[float | None]) -> str:
    numeric = [value for value in values if value is not None and value != 0]
    if not numeric:
        return "UNDETERMINED_ALL_ZERO_OR_MISSING"
    if all(value > 0 for value in numeric):
        return "POSITIVE_COST"
    if all(value < 0 for value in numeric):
        return "NEGATIVE_COST"
    return "MIXED_COST_SIGN"


def _normalized_cost_value(value: float | None) -> tuple[float | None, str]:
    if value is None:
        return None, "NOT_APPLICABLE_MISSING"
    if value < 0:
        return abs(value), "DISPLAY_NORMALIZED"
    return value, "NONE"


def _header_column(ws, label: str, *, min_col: int, max_col: int) -> int:
    found = [
        column
        for column in range(min_col, max_col + 1)
        if _matches(ws.cell(1, column).value, label)
    ]
    if len(found) != 1:
        raise ValueError(f"Expected one {label!r} header in columns {min_col}:{max_col}; found {found}")
    return found[0]


def _audit_actual_cost_source(wb) -> dict[str, Any]:
    """Validate the Actual Cost MTD category and sign contract without rebuilding it."""

    sheet_name = "2026 acutal cost"
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"Required sheet {sheet_name!r} not found")
    ws = wb[sheet_name]
    item_col = _header_column(ws, "Item", min_col=1, max_col=5)
    _header_column(ws, "Channel", min_col=1, max_col=5)
    _header_column(ws, "Brand", min_col=1, max_col=5)
    mtd_col = _header_column(ws, "MTD", min_col=6, max_col=21)
    label_to_name = {label: name for name, label in OPERATING_SOURCE_LABELS.items()}
    rows: list[dict[str, Any]] = []
    for row in range(2, ws.max_row + 1):
        label = _text(ws.cell(row, item_col).value)
        metric_name = label_to_name.get(label or "")
        if metric_name is None:
            continue
        cell = ws.cell(row, mtd_col)
        rows.append(
            {
                "metric_name": metric_name,
                "source_label": label,
                "channel": _text(ws.cell(row, item_col + 1).value),
                "brand": _text(ws.cell(row, item_col + 2).value),
                "source_reference": cell.coordinate,
                "source_value": as_number(cell.value),
            }
        )
    expected = set(OPERATING_SOURCE_LABELS)
    observed = {row["metric_name"] for row in rows}
    if observed != expected:
        raise ValueError(f"Actual Cost operating item coverage drift: {sorted(observed)}")
    cost_values = [
        row["source_value"] for row in rows if row["metric_name"] != "gmv"
    ]
    totals = {
        metric_name: sum(
            row["source_value"] or 0
            for row in rows
            if row["metric_name"] == metric_name
        )
        for metric_name in OPERATING_SOURCE_LABELS
    }
    return {
        "source_sheet": sheet_name,
        "period_field": _text(ws.cell(1, mtd_col).value),
        "period_source_reference": ws.cell(1, mtd_col).coordinate,
        "row_count": len(rows),
        "source_item_labels": list(OPERATING_SOURCE_LABELS.values()),
        "sign_convention": _cost_sign_convention(cost_values),
        "component_overlap_status": "NO_OVERLAP_OBSERVED_IN_SOURCE_CATEGORY_STRUCTURE",
        "component_overlap_basis": (
            "Each source row has one Item classification, and the management summary exposes "
            "MKT-Insite, Fixed Cost, Return+Warranty, and Funding as separate lines."
        ),
        "known_burden_allowed": True,
        "source_totals": totals,
        "rows": rows,
    }


def _record(
    *,
    snapshot_id: str,
    source_file: str,
    source_sha256: str,
    period: dict[str, Any],
    section: str,
    level: str,
    canonical: Any,
    raw_dimension: str,
    scenario: str,
    scenario_header: str,
    cell: Any,
) -> dict[str, Any]:
    raw_value = cell.value
    metric_value = as_number(raw_value)
    dimensions = _base_dimensions(level, canonical)
    quality = "VALID" if metric_value is not None else "DATA_QUALITY_WARNING"
    identity = ":".join(
        str(value or "NULL")
        for value in (
            level,
            dimensions["platform"],
            dimensions["channel"],
            dimensions["brand"],
            dimensions["power_source"],
            scenario,
            METRIC_NAME,
        )
    )
    return {
        "record_id": f"{snapshot_id}:CM:{identity}",
        "snapshot_id": snapshot_id,
        "period_label": period["period_label"],
        "period_start": None,
        "period_end": None,
        "data_through": None,
        **dimensions,
        "analysis_level": level,
        "metric_domain": METRIC_DOMAIN,
        "metric_name": METRIC_NAME if scenario != "ACTUAL_VS_BP" else "contribution_margin_rate_gap",
        "metric_value": metric_value,
        "unit": "ratio" if scenario != "ACTUAL_VS_BP" else "percentage_points",
        "value_scale": "fraction_of_one",
        "scenario": scenario,
        "metric_origin": "SOURCE_REPORTED",
        "source_file": source_file,
        "source_sha256": source_sha256,
        "source_sheet": "over view",
        "source_section": section,
        "source_header": f"{scenario_header} | {raw_dimension}",
        "source_reference": cell.coordinate,
        "mapping_status": "MAPPED",
        "data_quality_status": quality,
        "raw_value": raw_value,
        "raw_dimensions": {"source_dimension": raw_dimension},
    }


def _read_parallel_summary(
    ws,
    *,
    section_label: str,
    next_section_label: str,
    level: str,
    dimension_map: dict[str, Any],
    snapshot_id: str,
    source_file: str,
    source_sha256: str,
    period: dict[str, Any],
) -> list[dict[str, Any]]:
    anchor = _find_one(ws, section_label)
    next_anchor = _find_one(ws, next_section_label, min_row=anchor.row + 1)
    header_row = anchor.row + 1
    actual_headers = _find_in_row(ws, header_row, "US-Actual")
    bp_headers = _find_in_row(ws, header_row, "US-BP")
    gap_anchor = _find_one(ws, "Actual CM vs. BP CM", min_row=anchor.row, max_row=anchor.row)
    if len(actual_headers) != 1 or len(bp_headers) < 2:
        raise ValueError(f"Scenario header contract drift in {section_label}")
    actual_start = actual_headers[0].column
    bp_start = min(cell.column for cell in bp_headers if cell.column > actual_start)
    gap_start = gap_anchor.column
    if not _matches(ws.cell(header_row, gap_start).value, "US-BP") or not (actual_start < bp_start < gap_start):
        raise ValueError(f"Scenario order contract drift in {section_label}")
    groups = (
        ("ACTUAL", actual_start, bp_start - 1),
        ("BP", bp_start, gap_start - 1),
        ("ACTUAL_VS_BP", gap_start, ws.max_column),
    )
    records: list[dict[str, Any]] = []
    for scenario, start_col, end_col in groups:
        metric_row = _metric_row(ws, start_col, header_row, next_anchor.row - 1)
        scenario_header = _text(ws.cell(header_row, start_col).value) or "UNKNOWN"
        for column in range(start_col + 1, end_col + 1):
            raw_dimension = _text(ws.cell(header_row, column).value)
            canonical = dimension_map.get(raw_dimension or "")
            if canonical is None:
                continue
            records.append(
                _record(
                    snapshot_id=snapshot_id,
                    source_file=source_file,
                    source_sha256=source_sha256,
                    period=period,
                    section=section_label,
                    level=level,
                    canonical=canonical,
                    raw_dimension=raw_dimension or "UNKNOWN",
                    scenario=scenario,
                    scenario_header=scenario_header,
                    cell=ws.cell(metric_row, column),
                )
            )
    return records


def _first_label_in_column(ws, label: str, column: int, start_row: int, end_row: int) -> Any:
    found = [
        ws.cell(row, column)
        for row in range(start_row, end_row + 1)
        if _matches(ws.cell(row, column).value, label)
    ]
    if len(found) != 1:
        raise ValueError(f"Expected one {label!r} in column {column}, rows {start_row}:{end_row}")
    return found[0]


def _read_channel_summary(
    ws,
    *,
    snapshot_id: str,
    source_file: str,
    source_sha256: str,
    period: dict[str, Any],
) -> list[dict[str, Any]]:
    section = _find_one(ws, "All Brand")
    details = _find_one(ws, "Details", min_row=section.row + 1)
    actual = _first_label_in_column(ws, "US-Actual", section.column, section.row + 1, details.row - 1)
    gap_section = _find_one(
        ws, "Actual CM vs. BP CM", min_row=actual.row + 1, max_row=details.row - 1
    )
    bp = _first_label_in_column(ws, "US-BP", section.column, actual.row + 1, gap_section.row - 1)
    gap_header = gap_section.row + 1
    if not _matches(ws.cell(gap_header, section.column).value, "US-BP"):
        raise ValueError("All Brand gap scenario header contract drift")
    groups = (
        ("ACTUAL", actual.row, bp.row - 1),
        ("BP", bp.row, gap_section.row - 1),
        ("ACTUAL_VS_BP", gap_header, details.row - 1),
    )
    records: list[dict[str, Any]] = []
    for scenario, header_row, end_row in groups:
        start_col = section.column
        metric_row = _metric_row(ws, start_col, header_row, end_row)
        scenario_header = _text(ws.cell(header_row, start_col).value) or "UNKNOWN"
        seen: set[tuple[str, str]] = set()
        for column in range(start_col + 1, ws.max_column + 1):
            raw_dimension = _text(ws.cell(header_row, column).value)
            canonical = CHANNEL_MAP.get(raw_dimension or "")
            if canonical is None:
                continue
            if canonical in seen:
                raise ValueError(f"Duplicate All Brand channel header for {canonical} in {scenario}")
            seen.add(canonical)
            records.append(
                _record(
                    snapshot_id=snapshot_id,
                    source_file=source_file,
                    source_sha256=source_sha256,
                    period=period,
                    section="All Brand",
                    level="CHANNEL",
                    canonical=canonical,
                    raw_dimension=raw_dimension or "UNKNOWN",
                    scenario=scenario,
                    scenario_header=scenario_header,
                    cell=ws.cell(metric_row, column),
                )
            )
        if seen != set(CHANNEL_MAP.values()):
            raise ValueError(f"All Brand {scenario} channel coverage drift: {sorted(seen)}")
    return records


def _read_channel_gmv_context(
    ws,
    *,
    snapshot_id: str,
    source_file: str,
    source_sha256: str,
    period: dict[str, Any],
) -> list[dict[str, Any]]:
    """Read direct overview Actual GMV only to audit Sales/CM scope alignment."""

    section = _find_one(ws, "All Brand")
    details = _find_one(ws, "Details", min_row=section.row + 1)
    actual = _first_label_in_column(ws, "US-Actual", section.column, section.row + 1, details.row - 1)
    gap_section = _find_one(
        ws, "Actual CM vs. BP CM", min_row=actual.row + 1, max_row=details.row - 1
    )
    bp = _first_label_in_column(ws, "US-BP", section.column, actual.row + 1, gap_section.row - 1)
    gmv_row = _metric_row(ws, section.column, actual.row, bp.row - 1, metric="GMV")
    records: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for column in range(section.column + 1, ws.max_column + 1):
        raw_dimension = _text(ws.cell(actual.row, column).value)
        canonical = CHANNEL_MAP.get(raw_dimension or "")
        if canonical is None:
            continue
        if canonical in seen:
            raise ValueError(f"Duplicate All Brand Actual GMV header for {canonical}")
        seen.add(canonical)
        cell = ws.cell(gmv_row, column)
        value = as_number(cell.value)
        records.append(
            {
                "record_id": f"{snapshot_id}:CM:CHANNEL:{canonical[0]}:{canonical[1]}:ACTUAL:gmv",
                "snapshot_id": snapshot_id,
                "period_label": period["period_label"],
                "period_start": None,
                "period_end": None,
                "data_through": None,
                "platform": canonical[0],
                "channel": canonical[1],
                "subchannel": None,
                "brand": None,
                "power_source": None,
                "sku": None,
                "analysis_level": "CHANNEL",
                "metric_domain": METRIC_DOMAIN,
                "metric_name": "gmv",
                "metric_value": value,
                "unit": "USD",
                "value_scale": "absolute",
                "scenario": "ACTUAL",
                "metric_origin": "SOURCE_REPORTED",
                "source_file": source_file,
                "source_sha256": source_sha256,
                "source_sheet": "over view",
                "source_section": "All Brand",
                "source_header": f"US-Actual | {raw_dimension}",
                "source_reference": cell.coordinate,
                "mapping_status": "MAPPED",
                "data_quality_status": "VALID" if value is not None else "DATA_QUALITY_WARNING",
                "raw_value": cell.value,
                "raw_dimensions": {"source_dimension": raw_dimension},
            }
        )
    if seen != set(CHANNEL_MAP.values()):
        raise ValueError(f"All Brand Actual GMV channel coverage drift: {sorted(seen)}")
    return records


def _operating_source_record(
    *,
    snapshot_id: str,
    source_file: str,
    source_sha256: str,
    period: dict[str, Any],
    section: str,
    level: str,
    canonical: Any,
    raw_dimension: str,
    metric_name: str,
    source_label: str,
    scenario_header: str,
    cell: Any,
    cost_sign_convention: str,
) -> dict[str, Any]:
    source_value = as_number(cell.value)
    dimensions = _base_dimensions(level, canonical)
    is_cost = metric_name != "gmv"
    normalized, normalization = (
        _normalized_cost_value(source_value)
        if is_cost
        else (source_value, "NONE")
    )
    identity = ":".join(
        str(value or "NULL")
        for value in (
            level,
            dimensions["platform"],
            dimensions["channel"],
            dimensions["brand"],
            dimensions["power_source"],
            metric_name,
        )
    )
    return {
        "record_id": f"{snapshot_id}:OPERATING:{identity}",
        "snapshot_id": snapshot_id,
        "period_label": period["period_label"],
        "period_start": None,
        "period_end": None,
        "data_through": None,
        **dimensions,
        "analysis_level": level,
        "metric_domain": OPERATING_DOMAIN,
        "metric_name": metric_name,
        "metric_value": source_value,
        "unit": "USD",
        "value_scale": "absolute",
        "scenario": "ACTUAL",
        "metric_origin": "SOURCE_REPORTED",
        "source_value": source_value,
        "source_sign": _source_sign(source_value),
        "cost_sign_convention": cost_sign_convention if is_cost else "NOT_APPLICABLE",
        "normalized_display_value": normalized,
        "display_normalization": normalization,
        "source_file": source_file,
        "source_sha256": source_sha256,
        "source_sheet": "over view",
        "source_section": section,
        "source_label": source_label,
        "source_header": f"{scenario_header} | {raw_dimension}",
        "source_reference": cell.coordinate,
        "source_references": [cell.coordinate],
        "mapping_status": "MAPPED",
        "data_quality_status": "VALID" if source_value is not None else "DATA_QUALITY_WARNING",
        "derivation_status": "NOT_APPLICABLE_SOURCE_REPORTED",
        "formula": None,
        "source_values": None,
        "raw_value": cell.value,
        "raw_dimensions": {"source_dimension": raw_dimension},
    }


def _read_parallel_actual_operating(
    ws,
    *,
    section_label: str,
    next_section_label: str,
    level: str,
    dimension_map: dict[str, Any],
    snapshot_id: str,
    source_file: str,
    source_sha256: str,
    period: dict[str, Any],
    cost_sign_convention: str,
) -> list[dict[str, Any]]:
    anchor = _find_one(ws, section_label)
    next_anchor = _find_one(ws, next_section_label, min_row=anchor.row + 1)
    header_row = anchor.row + 1
    actual_headers = _find_in_row(ws, header_row, "US-Actual")
    bp_headers = _find_in_row(ws, header_row, "US-BP")
    if len(actual_headers) != 1 or len(bp_headers) < 2:
        raise ValueError(f"Scenario header contract drift in {section_label}")
    actual_start = actual_headers[0].column
    bp_start = min(cell.column for cell in bp_headers if cell.column > actual_start)
    metric_rows = {
        metric_name: _metric_row(
            ws,
            actual_start,
            header_row,
            next_anchor.row - 1,
            metric=source_label,
        )
        for metric_name, source_label in OPERATING_SOURCE_LABELS.items()
    }
    records: list[dict[str, Any]] = []
    for column in range(actual_start + 1, bp_start):
        raw_dimension = _text(ws.cell(header_row, column).value)
        canonical = dimension_map.get(raw_dimension or "")
        if canonical is None:
            continue
        for metric_name, source_label in OPERATING_SOURCE_LABELS.items():
            records.append(
                _operating_source_record(
                    snapshot_id=snapshot_id,
                    source_file=source_file,
                    source_sha256=source_sha256,
                    period=period,
                    section=section_label,
                    level=level,
                    canonical=canonical,
                    raw_dimension=raw_dimension or "UNKNOWN",
                    metric_name=metric_name,
                    source_label=source_label,
                    scenario_header="US-Actual",
                    cell=ws.cell(metric_rows[metric_name], column),
                    cost_sign_convention=cost_sign_convention,
                )
            )
    return records


def _read_channel_actual_operating(
    ws,
    *,
    snapshot_id: str,
    source_file: str,
    source_sha256: str,
    period: dict[str, Any],
    cost_sign_convention: str,
) -> list[dict[str, Any]]:
    section = _find_one(ws, "All Brand")
    details = _find_one(ws, "Details", min_row=section.row + 1)
    actual = _first_label_in_column(ws, "US-Actual", section.column, section.row + 1, details.row - 1)
    gap_section = _find_one(
        ws, "Actual CM vs. BP CM", min_row=actual.row + 1, max_row=details.row - 1
    )
    bp = _first_label_in_column(ws, "US-BP", section.column, actual.row + 1, gap_section.row - 1)
    metric_rows = {
        metric_name: _metric_row(
            ws,
            section.column,
            actual.row,
            bp.row - 1,
            metric=source_label,
        )
        for metric_name, source_label in OPERATING_SOURCE_LABELS.items()
    }
    records: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for column in range(section.column + 1, ws.max_column + 1):
        raw_dimension = _text(ws.cell(actual.row, column).value)
        canonical = CHANNEL_MAP.get(raw_dimension or "")
        if canonical is None:
            continue
        if canonical in seen:
            raise ValueError(f"Duplicate All Brand operating header for {canonical}")
        seen.add(canonical)
        for metric_name, source_label in OPERATING_SOURCE_LABELS.items():
            records.append(
                _operating_source_record(
                    snapshot_id=snapshot_id,
                    source_file=source_file,
                    source_sha256=source_sha256,
                    period=period,
                    section="All Brand",
                    level="CHANNEL",
                    canonical=canonical,
                    raw_dimension=raw_dimension or "UNKNOWN",
                    metric_name=metric_name,
                    source_label=source_label,
                    scenario_header="US-Actual",
                    cell=ws.cell(metric_rows[metric_name], column),
                    cost_sign_convention=cost_sign_convention,
                )
            )
    if seen != set(CHANNEL_MAP.values()):
        raise ValueError(f"All Brand operating channel coverage drift: {sorted(seen)}")
    return records


def _operating_key(record: dict[str, Any]) -> tuple[Any, ...]:
    level = record["analysis_level"]
    if level == "CHANNEL":
        return level, record["platform"], record["channel"]
    if level == "BRAND":
        return level, record["brand"]
    if level == "POWER_SOURCE":
        return level, record["power_source"]
    raise ValueError(f"Unsupported operating level: {level}")


def _derived_operating_record(
    template: dict[str, Any],
    *,
    metric_name: str,
    metric_value: Any,
    unit: str,
    formula: str,
    status: str,
    inputs: list[dict[str, Any]],
) -> dict[str, Any]:
    source_values = {row["metric_name"]: row.get("source_value") for row in inputs}
    source_references = [row["source_reference"] for row in inputs if row.get("source_reference")]
    return {
        **{key: template.get(key) for key in (
            "snapshot_id", "period_label", "period_start", "period_end", "data_through",
            "platform", "channel", "subchannel", "brand", "power_source", "sku",
            "analysis_level", "source_file", "source_sha256",
        )},
        "record_id": f"{template['snapshot_id']}:OPERATING_DERIVED:{':'.join(str(v or 'NULL') for v in (*_operating_key(template), metric_name))}",
        "metric_domain": OPERATING_DOMAIN,
        "metric_name": metric_name,
        "metric_value": metric_value,
        "unit": unit,
        "value_scale": "fraction_of_one" if unit == "ratio" else "absolute",
        "scenario": "ACTUAL",
        "metric_origin": "DERIVED",
        "source_value": None,
        "source_sign": "NOT_APPLICABLE",
        "cost_sign_convention": template.get("cost_sign_convention"),
        "normalized_display_value": metric_value,
        "display_normalization": "NOT_APPLICABLE_DERIVED",
        "source_sheet": "over view",
        "source_section": template.get("source_section"),
        "source_label": None,
        "source_header": None,
        "source_reference": None,
        "source_references": source_references,
        "mapping_status": template.get("mapping_status"),
        "data_quality_status": "VALID" if status == "DERIVED" else "NOT_AVAILABLE",
        "derivation_status": status,
        "formula": formula,
        "source_values": source_values,
        "raw_value": None,
        "raw_dimensions": template.get("raw_dimensions"),
    }


def derive_operating_efficiency_metrics(
    source_records: Iterable[dict[str, Any]],
    *,
    components_non_overlapping: bool,
) -> list[dict[str, Any]]:
    """Derive only the Phase 4 approved ratios, flag, and known-cost burden."""

    groups: dict[tuple[Any, ...], dict[str, dict[str, Any]]] = {}
    for record in source_records:
        if record.get("metric_origin") != "SOURCE_REPORTED":
            continue
        metrics = groups.setdefault(_operating_key(record), {})
        name = record["metric_name"]
        if name in metrics:
            raise ValueError(f"Duplicate operating source metric for {_operating_key(record)} / {name}")
        metrics[name] = record
    derived: list[dict[str, Any]] = []
    rate_specs = (
        ("tacos", "market_insight", "Market Insight / GMV"),
        ("fixed_cost_rate", "fixed_cost", "Fixed Cost / GMV"),
        (
            "return_warranty_cost_rate",
            "return_warranty_cost",
            "Return + Warranty / GMV",
        ),
        ("funding_rate", "funding", "Funding / GMV"),
    )
    for metrics in groups.values():
        template = next(iter(metrics.values()))
        gmv = metrics.get("gmv")
        gmv_value = gmv.get("source_value") if gmv else None
        for derived_name, numerator_name, formula in rate_specs:
            numerator = metrics.get(numerator_name)
            inputs = [row for row in (numerator, gmv) if row is not None]
            if numerator is None or numerator.get("normalized_display_value") is None or gmv_value is None:
                value, status = None, "NOT_DERIVED_MISSING_SOURCE"
            elif gmv_value <= 0:
                value, status = None, "NOT_DERIVED_ZERO_OR_NEGATIVE_GMV"
            else:
                value = numerator["normalized_display_value"] / gmv_value
                status = "DERIVED"
            derived.append(
                _derived_operating_record(
                    template,
                    metric_name=derived_name,
                    metric_value=value,
                    unit="ratio",
                    formula=formula,
                    status=status,
                    inputs=inputs,
                )
            )
        funding = metrics.get("funding")
        funding_value = funding.get("normalized_display_value") if funding else None
        derived.append(
            _derived_operating_record(
                template,
                metric_name="funding_present",
                metric_value=None if funding_value is None else ("YES" if funding_value > 0 else "NO"),
                unit="flag",
                formula="Funding > 0",
                status="DERIVED" if funding_value is not None else "NOT_DERIVED_MISSING_SOURCE",
                inputs=[funding] if funding else [],
            )
        )
        components = [metrics.get(name) for name in (
            "market_insight", "fixed_cost", "return_warranty_cost", "funding"
        )]
        available_components = [row for row in components if row is not None]
        if not components_non_overlapping:
            burden, burden_status = None, "NOT_DERIVED_DUE_TO_OVERLAPPING_COMPONENTS"
        elif len(available_components) != len(components) or any(
            row.get("normalized_display_value") is None for row in available_components
        ):
            burden, burden_status = None, "NOT_DERIVED_MISSING_SOURCE"
        else:
            burden = sum(row["normalized_display_value"] for row in available_components)
            burden_status = "DERIVED"
        derived.append(
            _derived_operating_record(
                template,
                metric_name="known_operating_cost_burden",
                metric_value=burden,
                unit="USD",
                formula="Market Insight + Fixed Cost + Return + Warranty + Funding",
                status=burden_status,
                inputs=available_components,
            )
        )
        burden_inputs = [*available_components, *([gmv] if gmv else [])]
        if burden is None:
            burden_rate, burden_rate_status = None, burden_status
        elif gmv_value is None:
            burden_rate, burden_rate_status = None, "NOT_DERIVED_MISSING_SOURCE"
        elif gmv_value <= 0:
            burden_rate, burden_rate_status = None, "NOT_DERIVED_ZERO_OR_NEGATIVE_GMV"
        else:
            burden_rate, burden_rate_status = burden / gmv_value, "DERIVED"
        derived.append(
            _derived_operating_record(
                template,
                metric_name="known_operating_cost_burden_rate",
                metric_value=burden_rate,
                unit="ratio",
                formula="Known Operating Cost Burden / GMV",
                status=burden_rate_status,
                inputs=burden_inputs,
            )
        )
    return derived


def _operating_coverage(records: list[dict[str, Any]]) -> dict[str, Any]:
    source = [row for row in records if row["metric_origin"] == "SOURCE_REPORTED"]
    derived = [row for row in records if row["metric_origin"] == "DERIVED"]
    levels = ("CHANNEL", "BRAND", "POWER_SOURCE")
    return {
        "record_count": len(records),
        "source_reported_record_count": len(source),
        "derived_record_count": len(derived),
        "source_by_level": {
            level: sum(row["analysis_level"] == level for row in source) for level in levels
        },
        "derived_by_level": {
            level: sum(row["analysis_level"] == level for row in derived) for level in levels
        },
        "source_metrics": sorted({row["metric_name"] for row in source}),
        "derived_metrics": sorted({row["metric_name"] for row in derived}),
        "sku_operating_metrics": 0,
    }


def _reconcile_operating(
    records: list[dict[str, Any]],
    source_audit: dict[str, Any],
) -> dict[str, Any]:
    source = [row for row in records if row["metric_origin"] == "SOURCE_REPORTED"]
    derived = [row for row in records if row["metric_origin"] == "DERIVED"]
    overview_brand_totals = {
        metric_name: sum(
            row["source_value"] or 0
            for row in source
            if row["analysis_level"] == "BRAND" and row["metric_name"] == metric_name
        )
        for metric_name in OPERATING_SOURCE_LABELS
    }
    total_reconciliation = {
        metric_name: {
            "actual_cost_mtd": source_audit["source_totals"][metric_name],
            "overview_brand_total": overview_brand_totals[metric_name],
            "delta": overview_brand_totals[metric_name] - source_audit["source_totals"][metric_name],
        }
        for metric_name in OPERATING_SOURCE_LABELS
    }
    identities = [(_operating_key(row), row["metric_name"], row["metric_origin"]) for row in records]
    def formula_valid(row: dict[str, Any]) -> bool:
        if row["derivation_status"] != "DERIVED":
            return row["metric_value"] is None
        values = row["source_values"] or {}
        gmv = values.get("gmv")
        magnitudes = {
            name: abs(value) if value is not None else None
            for name, value in values.items()
            if name != "gmv"
        }
        name = row["metric_name"]
        if name == "funding_present":
            expected = "YES" if (magnitudes.get("funding") or 0) > 0 else "NO"
            return row["metric_value"] == expected
        if name == "known_operating_cost_burden":
            expected = sum(magnitudes[component] for component in (
                "market_insight", "fixed_cost", "return_warranty_cost", "funding"
            ))
        elif name == "known_operating_cost_burden_rate":
            expected = sum(magnitudes[component] for component in (
                "market_insight", "fixed_cost", "return_warranty_cost", "funding"
            )) / gmv
        else:
            numerator = {
                "tacos": "market_insight",
                "fixed_cost_rate": "fixed_cost",
                "return_warranty_cost_rate": "return_warranty_cost",
                "funding_rate": "funding",
            }[name]
            expected = magnitudes[numerator] / gmv
        return abs(row["metric_value"] - expected) <= 1e-12

    checks = {
        "source_metric_identity": len(source) == 50,
        "derived_metric_identity": len(derived) == 70,
        "no_duplicate_reading": len(identities) == len(set(identities)),
        "source_reported_values_not_mutated": all(
            row["metric_value"] == row["source_value"] == row["raw_value"] for row in source
        ),
        "actual_cost_total_reconciliation": all(
            abs(item["delta"]) <= 1e-6 for item in total_reconciliation.values()
        ),
        "sign_convention_verified": source_audit["sign_convention"] in {
            "POSITIVE_COST", "NEGATIVE_COST"
        },
        "no_sku_or_lower_level_allocation": all(
            row["analysis_level"] != "SKU" and row["sku"] is None for row in records
        ),
        "derived_metrics_traceable": all(
            row["source_references"] and row["formula"] for row in derived
        ),
        "derived_value_formula_reconciliation": all(formula_valid(row) for row in derived),
        "known_burden_not_cm": all(
            row["metric_name"] not in {"contribution_margin", "contribution_margin_rate"}
            for row in derived
        ),
    }
    return {
        "status": "PASS" if all(checks.values()) else "FAIL",
        "checks": checks,
        "source_total_reconciliation": total_reconciliation,
        "cm_recalculation": "NONE",
    }


def _coverage(records: Iterable[dict[str, Any]]) -> dict[str, Any]:
    rows = list(records)
    cm_rows = [row for row in rows if row["metric_name"] != "gmv"]
    levels = ("CHANNEL", "BRAND", "POWER_SOURCE")
    scenarios = ("ACTUAL", "BP", "ACTUAL_VS_BP")
    return {
        "record_count": len(cm_rows),
        "business_record_count": len(rows),
        "channel_actual_gmv_context_count": sum(row["metric_name"] == "gmv" for row in rows),
        "by_level": {level: sum(row["analysis_level"] == level for row in cm_rows) for level in levels},
        "by_scenario": {scenario: sum(row["scenario"] == scenario for row in cm_rows) for scenario in scenarios},
        "actual_cm": any(row["scenario"] == "ACTUAL" for row in cm_rows),
        "bp_cm": any(row["scenario"] == "BP" for row in cm_rows),
        "gap": any(row["scenario"] == "ACTUAL_VS_BP" for row in cm_rows),
    }


def _reconcile(records: list[dict[str, Any]]) -> dict[str, Any]:
    cm_records = [row for row in records if row["metric_name"] != "gmv"]
    identities = [
        (
            row["analysis_level"], row["platform"], row["channel"], row["brand"],
            row["power_source"], row["metric_name"], row["scenario"],
        )
        for row in records
    ]
    source_integrity = all(
        row["metric_value"] == row["raw_value"]
        for row in records
        if row["metric_value"] is not None
    )
    level_integrity = all(
        (
            row["analysis_level"] == "CHANNEL"
            and row["platform"] is not None
            and row["channel"] is not None
            and row["brand"] is None
            and row["power_source"] is None
            and row["sku"] is None
        )
        or (
            row["analysis_level"] == "BRAND"
            and row["brand"] is not None
            and row["platform"] is None
            and row["channel"] is None
            and row["power_source"] is None
            and row["sku"] is None
        )
        or (
            row["analysis_level"] == "POWER_SOURCE"
            and row["power_source"] is not None
            and row["platform"] is None
            and row["channel"] is None
            and row["brand"] is None
            and row["sku"] is None
        )
        for row in records
    )
    checks = {
        "metric_identity": len(cm_records) == 30 and len(records) == 34,
        "no_duplicate_reading": len(identities) == len(set(identities)),
        "level_integrity": level_integrity,
        "source_reported_integrity": source_integrity,
        "metric_origin_integrity": all(row["metric_origin"] == "SOURCE_REPORTED" for row in records),
        "cm_time_independent": all(row["data_through"] is None for row in records),
    }
    return {
        "status": "PASS" if all(checks.values()) else "FAIL",
        "checks": checks,
        "derived_metric_audit": [],
    }


def read_business_performance(
    workbook_path: Path,
    project_root: Path | None = None,
    snapshot_date: date | None = None,
) -> BusinessPerformanceData:
    """Locate, read, normalize, validate, and expose overview CM rates."""

    workbook_path = workbook_path.resolve()
    project_root = (project_root or Path(__file__).resolve().parents[2]).resolve()
    del project_root  # Reserved for future approved mapping-version lookup.
    snapshot_date = snapshot_date or date.today()
    source_hash = _sha256(workbook_path)
    snapshot_id = f"{snapshot_date.isoformat()}_{source_hash[:12].lower()}"
    wb = load_workbook(workbook_path, read_only=False, data_only=True, keep_links=True)
    if "over view" not in wb.sheetnames:
        wb.close()
        raise ValueError("Required sheet 'over view' not found")
    ws = wb["over view"]
    month = _text(ws["B2"].value)
    basis = _text(ws["B3"].value)
    if not month or not basis:
        wb.close()
        raise ValueError("Overview period label contract drift: B2/B3 labels are required")
    period = {
        "period_label": f"{month} {basis}",
        "month_label": month,
        "period_basis": basis,
        "period_start": None,
        "period_end": None,
        "data_through": None,
        "freshness_status": "UNKNOWN_DAY_LEVEL_DATA_THROUGH",
        "source_references": ["B2", "B3"],
    }
    source_audit = _audit_actual_cost_source(wb)
    records: list[dict[str, Any]] = []
    records.extend(
        _read_parallel_summary(
            ws,
            section_label="By Brand",
            next_section_label="By Power Source",
            level="BRAND",
            dimension_map={"Sunseeker": "Sunseeker", "Badger": "Badger"},
            snapshot_id=snapshot_id,
            source_file=workbook_path.name,
            source_sha256=source_hash,
            period=period,
        )
    )
    records.extend(
        _read_parallel_summary(
            ws,
            section_label="By Power Source",
            next_section_label="By Channel",
            level="POWER_SOURCE",
            dimension_map={item: item for item in ("Robot", "Gas", "Lithium", "ACC")},
            snapshot_id=snapshot_id,
            source_file=workbook_path.name,
            source_sha256=source_hash,
            period=period,
        )
    )
    records.extend(
        _read_channel_summary(
            ws,
            snapshot_id=snapshot_id,
            source_file=workbook_path.name,
            source_sha256=source_hash,
            period=period,
        )
    )
    records.extend(
        _read_channel_gmv_context(
            ws,
            snapshot_id=snapshot_id,
            source_file=workbook_path.name,
            source_sha256=source_hash,
            period=period,
        )
    )
    operating_source_records: list[dict[str, Any]] = []
    operating_source_records.extend(
        _read_parallel_actual_operating(
            ws,
            section_label="By Brand",
            next_section_label="By Power Source",
            level="BRAND",
            dimension_map={"Sunseeker": "Sunseeker", "Badger": "Badger"},
            snapshot_id=snapshot_id,
            source_file=workbook_path.name,
            source_sha256=source_hash,
            period=period,
            cost_sign_convention=source_audit["sign_convention"],
        )
    )
    operating_source_records.extend(
        _read_parallel_actual_operating(
            ws,
            section_label="By Power Source",
            next_section_label="By Channel",
            level="POWER_SOURCE",
            dimension_map={item: item for item in ("Robot", "Gas", "Lithium", "ACC")},
            snapshot_id=snapshot_id,
            source_file=workbook_path.name,
            source_sha256=source_hash,
            period=period,
            cost_sign_convention=source_audit["sign_convention"],
        )
    )
    operating_source_records.extend(
        _read_channel_actual_operating(
            ws,
            snapshot_id=snapshot_id,
            source_file=workbook_path.name,
            source_sha256=source_hash,
            period=period,
            cost_sign_convention=source_audit["sign_convention"],
        )
    )
    operating_records = operating_source_records + derive_operating_efficiency_metrics(
        operating_source_records,
        components_non_overlapping=source_audit["known_burden_allowed"],
    )
    operating_reconciliation = _reconcile_operating(operating_records, source_audit)
    technical_errors = [
        cell.coordinate
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("#")
    ]
    wb.close()
    quality_issues = [
        {
            "code": "CORE_CM_SOURCE_ERROR",
            "status": "DATA_QUALITY_WARNING",
            "summary": "A required source-reported CM rate is an Excel error or non-numeric value.",
            "affected_records": sum(row["metric_value"] is None for row in records),
            "source_references": [
                row["source_reference"] for row in records if row["metric_value"] is None
            ],
            "affects_business_judgment": True,
        }
    ] if any(row["metric_value"] is None for row in records) else []
    if source_audit["sign_convention"] == "MIXED_COST_SIGN":
        quality_issues.append(
            {
                "code": "MIXED_OPERATING_COST_SIGN",
                "status": "DATA_QUALITY_WARNING",
                "summary": "Operating cost source values use mixed signs; cost rates are not decision-ready.",
                "affects_business_judgment": True,
            }
        )
    if operating_reconciliation["status"] != "PASS":
        quality_issues.append(
            {
                "code": "OPERATING_SOURCE_RECONCILIATION_FAILED",
                "status": "DATA_QUALITY_WARNING",
                "summary": "Operating source or derived metric reconciliation did not pass.",
                "affects_business_judgment": True,
            }
        )
    source = {
        "source_file": workbook_path.name,
        "source_sha256": source_hash,
        "source_sheet": "over view",
        "parser_version": PHASE3_CONTRACT_VERSION,
        "operating_parser_version": PHASE4_CONTRACT_VERSION,
        "technical_overview_error_count": len(technical_errors),
        "technical_overview_error_sample": technical_errors[:10],
    }
    return BusinessPerformanceData(
        snapshot_id=snapshot_id,
        source=source,
        period=period,
        records=records,
        coverage=_coverage(records),
        quality_issues=quality_issues,
        reconciliation=_reconcile(records),
        operating_records=operating_records,
        operating_coverage=_operating_coverage(operating_records),
        operating_reconciliation=operating_reconciliation,
        operating_source_audit=source_audit,
    )

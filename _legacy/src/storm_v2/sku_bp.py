"""Phase 5R SKU BP extraction and Actual x BP integration.

The authoritative target source is KPI Rawdata in the approved CM workbook.
The source workbook is opened read-only and never modified. Actual Sales stays
on the frozen STORM V2 Phase 2 sales-fact snapshot.
"""

from __future__ import annotations

import calendar
import csv
import hashlib
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import load_workbook

from .normalization import as_number, normalize_sku, normalize_text
from .sales_engine import aggregate_sales


PHASE5R_DATA_STATUS = "STORM_V2_PHASE_5R_DATA_READY_VISUAL_PENDING"
PHASE5R_READY_STATUS = "STORM_V2_PHASE_5R_WEEKLY_COCKPIT_READY"
BP_SHEET = "KPI Rawdata"
PROMOTION_HEADER_ALIASES = ("Promotion Type", "Promition Type")
CUSTOMER_CHANNELS = {
    "The Home Depot Inc": ("THD", "DS"),
    "Lowe's": ("Lowe's", "DS"),
    "Walmart DSV": ("Walmart", "DSV"),
    "Walmart Seller": ("Walmart", "MP"),
}
CUSTOMER_PLANNER_SHEETS = {
    "The Home Depot Inc": "The Home Depot-WBP",
    "Lowe's": "Lowes-WBP",
    "Walmart DSV": "Walmart DSV-WBP",
    "Walmart Seller": "Walmart Seller-WBP",
}
# The official Walmart Seller planning sheet defines the marketplace SKU scope.
# KPI Rawdata also contains a Sunseeker Robot promotion block for that customer;
# it is not present in the planner roster and must not enter the Badger WBP rollup.
PLANNER_SKU_SCOPE_CUSTOMERS = {"Walmart Seller"}
BP_GROUP_FIELDS = ("platform", "channel", "brand", "power_source", "sku")


class SkuBpContractError(ValueError):
    """Raised when the approved SKU BP source contract no longer holds."""


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def _header_index(sheet: Any) -> dict[str, int]:
    headers = [normalize_text(cell.value) for cell in sheet[1]]
    return {name: index for index, name in enumerate(headers) if name}


def _promotion_header(index: dict[str, int]) -> str:
    for name in PROMOTION_HEADER_ALIASES:
        if name in index:
            return name
    raise SkuBpContractError(
        "KPI Rawdata missing Promotion Type/Promition Type source component"
    )


def _month_label(month: int) -> str:
    return calendar.month_abbr[month]


def _key(row: dict[str, Any], fields: Sequence[str]) -> tuple[Any, ...]:
    return tuple(row.get(field) for field in fields)


def extract_sku_bp(
    workbook_path: Path,
    *,
    year: int,
    month: int,
    badger_gas_control: float | None = None,
) -> dict[str, Any]:
    """Extract monthly BP by Channel x SKU while summing every source line."""

    workbook_path = workbook_path.resolve()
    if not workbook_path.exists():
        raise FileNotFoundError(workbook_path)
    hash_before = _sha256(workbook_path)
    workbook = load_workbook(
        workbook_path, read_only=True, data_only=True, keep_links=False
    )
    if BP_SHEET not in workbook.sheetnames:
        workbook.close()
        raise SkuBpContractError(f"{BP_SHEET} sheet is missing")
    sheet = workbook[BP_SHEET]
    index = _header_index(sheet)
    required = {
        "SKU",
        "Customer",
        "Month-INT",
        "Year",
        "Month",
        "Brand",
        "Power Source",
        "TTL amount",
    }
    missing = sorted(required - set(index))
    if missing:
        workbook.close()
        raise SkuBpContractError(f"KPI Rawdata missing required headers: {missing}")
    promotion_header = _promotion_header(index)

    planner_sku_rosters: dict[str, set[str]] = {}
    for customer in PLANNER_SKU_SCOPE_CUSTOMERS:
        planner_sheet = CUSTOMER_PLANNER_SHEETS[customer]
        if planner_sheet not in workbook.sheetnames:
            workbook.close()
            raise SkuBpContractError(f"Planner sheet {planner_sheet!r} is missing")
        planner = workbook[planner_sheet]
        roster = {
            normalize_sku(values[0])
            for values in planner.iter_rows(min_row=8, values_only=True)
            if values and normalize_sku(values[0])
        }
        if not roster:
            workbook.close()
            raise SkuBpContractError(f"Planner sheet {planner_sheet!r} has no SKU roster")
        planner_sku_rosters[customer] = roster

    aggregate: dict[tuple[Any, ...], dict[str, Any]] = {}
    channel_sku_dimensions: dict[tuple[str, str, str], tuple[str, str]] = {}
    raw_channel: defaultdict[tuple[str, str], float] = defaultdict(float)
    raw_brand_power: defaultdict[tuple[str, str, str, str], float] = defaultdict(float)
    scoped_source_rows = 0
    blank_amount_rows: list[int] = []
    excluded_source_rows: list[dict[str, Any]] = []
    expected_month_label = _month_label(month)

    for row_number, values in enumerate(
        sheet.iter_rows(min_row=2, values_only=True), 2
    ):
        customer = normalize_text(values[index["Customer"]])
        if customer not in CUSTOMER_CHANNELS:
            continue
        source_year = as_number(values[index["Year"]])
        source_month = as_number(values[index["Month-INT"]])
        if source_year != year or source_month != month:
            continue
        source_month_label = normalize_text(values[index["Month"]])
        if source_month_label != expected_month_label:
            workbook.close()
            raise SkuBpContractError(
                f"row {row_number} Month={source_month_label!r}, "
                f"expected {expected_month_label!r}"
            )
        platform, channel = CUSTOMER_CHANNELS[customer]
        sku = normalize_sku(values[index["SKU"]])
        brand = normalize_text(values[index["Brand"]])
        power_source = normalize_text(values[index["Power Source"]])
        if not all((sku, brand, power_source)):
            workbook.close()
            raise SkuBpContractError(
                f"row {row_number} lacks SKU, Brand, or Power Source"
            )
        amount = as_number(values[index["TTL amount"]])
        if amount is None:
            blank_amount_rows.append(row_number)
            continue
        amount = float(amount)
        if customer in PLANNER_SKU_SCOPE_CUSTOMERS and sku not in planner_sku_rosters[customer]:
            excluded_source_rows.append(
                {
                    "source_sheet": BP_SHEET,
                    "source_row": row_number,
                    "logical_key": f"{customer}|{sku}|{brand}|{power_source}",
                    "customer": customer,
                    "platform": platform,
                    "channel": channel,
                    "brand": brand,
                    "power_source": power_source,
                    "sku": sku,
                    "bp_sales": amount,
                    "reason": "SKU_NOT_IN_OFFICIAL_WALMART_SELLER_WBP_ROSTER",
                    "scope_source_sheet": CUSTOMER_PLANNER_SHEETS[customer],
                }
            )
            continue
        promotion_type = normalize_text(values[index[promotion_header]]) or "UNKNOWN"
        dimension_key = (platform, channel, sku)
        observed = (brand, power_source)
        previous = channel_sku_dimensions.setdefault(dimension_key, observed)
        if previous != observed:
            workbook.close()
            raise SkuBpContractError(
                f"conflicting Brand/Power Source for {dimension_key}: "
                f"{previous!r} versus {observed!r}"
            )

        key = (platform, channel, brand, power_source, sku)
        bucket = aggregate.setdefault(
            key,
            {
                "bp_sales": 0.0,
                "source_rows": [],
                "promotion_components": defaultdict(float),
                "customers": set(),
            },
        )
        bucket["bp_sales"] += amount
        bucket["source_rows"].append(row_number)
        bucket["promotion_components"][promotion_type] += amount
        bucket["customers"].add(customer)
        raw_channel[(platform, channel)] += amount
        raw_brand_power[(platform, channel, brand, power_source)] += amount
        scoped_source_rows += 1

    workbook.close()
    hash_after = _sha256(workbook_path)
    if hash_after != hash_before:
        raise SkuBpContractError("source workbook changed during read-only extraction")

    rows = []
    for key, bucket in sorted(aggregate.items()):
        platform, channel, brand, power_source, sku = key
        rows.append(
            {
                "year": year,
                "month": month,
                "month_label": expected_month_label,
                "platform": platform,
                "channel": channel,
                "brand": brand,
                "power_source": power_source,
                "sku": sku,
                "bp_sales": bucket["bp_sales"],
                "promotion_types": sorted(bucket["promotion_components"]),
                "promotion_components": dict(
                    sorted(bucket["promotion_components"].items())
                ),
                "source_customers": sorted(bucket["customers"]),
                "source_rows": bucket["source_rows"],
                "source_sheet": BP_SHEET,
                "source_header": "TTL amount",
            }
        )

    control_actual = sum(
        row["bp_sales"]
        for row in rows
        if row["platform"] == "Walmart"
        and row["channel"] == "MP"
        and row["brand"] == "Badger"
        and row["power_source"] == "Gas"
    )
    return {
        "source": {
            "path": str(workbook_path),
            "file": workbook_path.name,
            "sha256": hash_before,
            "sheet": BP_SHEET,
            "promotion_header": promotion_header,
            "workbook_mutation": "NONE",
        },
        "period": {
            "year": year,
            "month": month,
            "month_label": expected_month_label,
            "period_start": f"{year:04d}-{month:02d}-01",
            "period_end": (
                f"{year:04d}-{month:02d}-"
                f"{calendar.monthrange(year, month)[1]:02d}"
            ),
        },
        "rows": rows,
        "source_row_count": scoped_source_rows,
        "blank_amount_rows": blank_amount_rows,
        "scope_definition": {
            "rule": "Walmart Seller KPI Rawdata is admitted only when SKU appears in the official Walmart Seller-WBP roster",
            "planner_sheets": {
                customer: {
                    "sheet": CUSTOMER_PLANNER_SHEETS[customer],
                    "sku_count": len(roster),
                }
                for customer, roster in planner_sku_rosters.items()
            },
        },
        "excluded_source_rows": excluded_source_rows,
        "excluded_bp_sales": sum(item["bp_sales"] for item in excluded_source_rows),
        "raw_rollups": {
            "channel": [
                {
                    "platform": key[0],
                    "channel": key[1],
                    "bp_sales": value,
                }
                for key, value in sorted(raw_channel.items())
            ],
            "brand_power": [
                {
                    "platform": key[0],
                    "channel": key[1],
                    "brand": key[2],
                    "power_source": key[3],
                    "bp_sales": value,
                }
                for key, value in sorted(raw_brand_power.items())
            ],
        },
        "walmart_badger_gas_control": {
            "expected": badger_gas_control,
            "actual": control_actual,
            "status": (
                "PASS"
                if badger_gas_control is not None and abs(control_actual - badger_gas_control) < 0.005
                else "SKIP_NO_BASELINE"
                if badger_gas_control is None
                else "FAIL"
            ),
        },
    }


def aggregate_bp(
    rows: Iterable[dict[str, Any]], group_by: Sequence[str]
) -> list[dict[str, Any]]:
    group_by = tuple(group_by)
    if any(field not in BP_GROUP_FIELDS for field in group_by):
        raise ValueError(f"group_by must be drawn from {BP_GROUP_FIELDS}")
    buckets: dict[tuple[Any, ...], dict[str, Any]] = {}
    for row in rows:
        key = _key(row, group_by)
        bucket = buckets.setdefault(
            key, {"bp_sales": 0.0, "source_rows": set(), "promotion_types": set()}
        )
        bucket["bp_sales"] += float(row["bp_sales"])
        bucket["source_rows"].update(row.get("source_rows", []))
        bucket["promotion_types"].update(row.get("promotion_types", []))
    output = []
    for key, bucket in sorted(buckets.items()):
        item = {field: value for field, value in zip(group_by, key)}
        item.update(
            {
                "bp_sales": bucket["bp_sales"],
                "source_rows": sorted(bucket["source_rows"]),
                "promotion_types": sorted(bucket["promotion_types"]),
            }
        )
        output.append(item)
    return output


def load_actual_sku_sales(snapshot_csv: Path) -> list[dict[str, Any]]:
    """Aggregate the frozen Actual Sales snapshot at full business/SKU grain."""

    with snapshot_csv.open(encoding="utf-8", newline="") as handle:
        facts = list(csv.DictReader(handle))
    return aggregate_sales(facts, BP_GROUP_FIELDS)


def _unique_index(
    rows: Iterable[dict[str, Any]],
    *,
    key_fields: Sequence[str],
    label: str,
) -> dict[tuple[Any, ...], dict[str, Any]]:
    output: dict[tuple[Any, ...], dict[str, Any]] = {}
    for row in rows:
        key = _key(row, key_fields)
        if key in output:
            raise SkuBpContractError(f"duplicate {label} key: {key}")
        output[key] = row
    return output


def join_actual_bp(
    actual_rows: Iterable[dict[str, Any]],
    bp_rows: Iterable[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Outer join on Channel + SKU + period, retaining zero-actual BP SKUs."""

    key_fields = ("platform", "channel", "sku")
    actual_index = _unique_index(actual_rows, key_fields=key_fields, label="Actual")
    bp_index = _unique_index(bp_rows, key_fields=key_fields, label="BP")
    joined = []
    for key in sorted(set(actual_index) | set(bp_index)):
        actual_row = actual_index.get(key)
        bp_row = bp_index.get(key)
        actual_sales = (
            float(actual_row["actual_sales"])
            if actual_row and actual_row.get("actual_sales") is not None
            else 0.0
        )
        bp_sales = (
            float(bp_row["bp_sales"])
            if bp_row and bp_row.get("bp_sales") is not None
            else None
        )
        actual_brand = normalize_text(actual_row.get("brand")) if actual_row else None
        actual_power = (
            normalize_text(actual_row.get("power_source")) if actual_row else None
        )
        bp_brand = normalize_text(bp_row.get("brand")) if bp_row else None
        bp_power = normalize_text(bp_row.get("power_source")) if bp_row else None
        if (
            actual_brand not in (None, "UNKNOWN")
            and bp_brand not in (None, "UNKNOWN")
            and actual_brand != bp_brand
        ):
            raise SkuBpContractError(
                f"Brand mismatch for {key}: Actual={actual_brand}, BP={bp_brand}"
            )
        if (
            actual_power not in (None, "UNKNOWN")
            and bp_power not in (None, "UNKNOWN")
            and actual_power != bp_power
        ):
            raise SkuBpContractError(
                f"Power Source mismatch for {key}: "
                f"Actual={actual_power}, BP={bp_power}"
            )
        gap = actual_sales - bp_sales if bp_sales is not None else None
        if bp_sales is None:
            attainment = None
            bp_status = "BP_NOT_AVAILABLE"
        elif bp_sales > 0:
            attainment = actual_sales / bp_sales
            bp_status = "MATCHED"
        elif bp_sales == 0 and actual_sales > 0:
            attainment = None
            bp_status = "BP_ZERO_ATTAINMENT_NA"
        else:
            attainment = None
            bp_status = "BOTH_ZERO_ATTAINMENT_NA"
        joined.append(
            {
                "platform": key[0],
                "channel": key[1],
                "brand": bp_brand or actual_brand or "UNKNOWN",
                "power_source": bp_power or actual_power or "UNKNOWN",
                "sku": key[2],
                "actual_sales": actual_sales,
                "bp_sales": bp_sales,
                "sales_gap": gap,
                "attainment": attainment,
                "bp_status": bp_status,
                "actual_present": actual_row is not None,
                "bp_present": bp_row is not None,
                "promotion_types": bp_row.get("promotion_types", []) if bp_row else [],
                "bp_source_rows": bp_row.get("source_rows", []) if bp_row else [],
                "actual_source_references": (
                    actual_row.get("source_references", []) if actual_row else []
                ),
            }
        )
    total_actual = sum(row["actual_sales"] for row in joined)
    for row in joined:
        row["contribution"] = (
            row["actual_sales"] / total_actual if total_actual else None
        )
    return joined


def gap_rankings(
    rows: Iterable[dict[str, Any]], *, limit: int = 8
) -> dict[str, list[dict[str, Any]]]:
    matched = [row for row in rows if row.get("sales_gap") is not None]
    overperformers = [
        row
        for row in matched
        if row.get("bp_sales") is not None
        and float(row["bp_sales"]) > 0
        and float(row["actual_sales"]) > float(row["bp_sales"])
    ]
    unplanned = [
        row
        for row in rows
        if (row.get("bp_sales") is None or float(row["bp_sales"]) == 0)
        and float(row["actual_sales"]) > 0
    ]
    return {
        "top_detractors": sorted(
            matched, key=lambda row: float(row["sales_gap"])
        )[:limit],
        "top_overperformers": sorted(
            overperformers, key=lambda row: float(row["sales_gap"]), reverse=True
        )[:limit],
        "unplanned_sales": sorted(
            unplanned, key=lambda row: float(row["actual_sales"]), reverse=True
        )[:limit],
    }


def reconcile_bp(extraction: dict[str, Any], *, scope_exclusion_expected: float | None = None) -> dict[str, Any]:
    """Reconcile the admitted official scope from SKU through TOTAL."""

    rows = extraction["rows"]
    derived_brand_power = aggregate_bp(
        rows, ("platform", "channel", "brand", "power_source")
    )
    derived_channel = aggregate_bp(rows, ("platform", "channel"))
    raw_brand_power = {
        _key(row, ("platform", "channel", "brand", "power_source")): row[
            "bp_sales"
        ]
        for row in extraction["raw_rollups"]["brand_power"]
    }
    raw_channel = {
        _key(row, ("platform", "channel")): row["bp_sales"]
        for row in extraction["raw_rollups"]["channel"]
    }
    level1 = [
        {
            **{field: row[field] for field in ("platform", "channel", "brand", "power_source")},
            "sku_rollup": row["bp_sales"],
            "raw_parent": raw_brand_power[
                _key(row, ("platform", "channel", "brand", "power_source"))
            ],
            "difference": row["bp_sales"]
            - raw_brand_power[
                _key(row, ("platform", "channel", "brand", "power_source"))
            ],
        }
        for row in derived_brand_power
    ]
    level2 = [
        {
            "platform": row["platform"],
            "channel": row["channel"],
            "brand_power_rollup": sum(
                item["bp_sales"]
                for item in derived_brand_power
                if item["platform"] == row["platform"]
                and item["channel"] == row["channel"]
            ),
            "sku_rollup": row["bp_sales"],
        }
        for row in derived_channel
    ]
    for item in level2:
        item["difference"] = item["brand_power_rollup"] - item["sku_rollup"]
    level3 = [
        {
            "platform": row["platform"],
            "channel": row["channel"],
            "derived_parent": row["bp_sales"],
            "raw_parent": raw_channel[_key(row, ("platform", "channel"))],
            "difference": row["bp_sales"]
            - raw_channel[_key(row, ("platform", "channel"))],
        }
        for row in derived_channel
    ]
    platform_rows = aggregate_bp(rows, ("platform",))
    level4 = []
    for platform in platform_rows:
        channel_total = sum(
            row["bp_sales"]
            for row in derived_channel
            if row["platform"] == platform["platform"]
        )
        level4.append(
            {
                "platform": platform["platform"],
                "channel_rollup": channel_total,
                "platform_rollup": platform["bp_sales"],
                "difference": channel_total - platform["bp_sales"],
            }
        )
    total_from_platform = sum(row["bp_sales"] for row in platform_rows)
    total_from_sku = sum(row["bp_sales"] for row in rows)
    total = {
        "platform_rollup": total_from_platform,
        "sku_rollup": total_from_sku,
        "difference": total_from_platform - total_from_sku,
    }
    status = (
        "PASS"
        if all(
            abs(item["difference"]) < 1e-9
            for item in (*level1, *level2, *level3, *level4, total)
        )
        and extraction["walmart_badger_gas_control"]["status"] in ("PASS", "SKIP_NO_BASELINE")
        and (
            scope_exclusion_expected is None
            or abs(extraction["excluded_bp_sales"] - scope_exclusion_expected) < 0.005
        )
        else "FAIL"
    )
    return {
        "status": status,
        "level_1_sku_to_brand_power": level1,
        "level_2_brand_power_to_channel": level2,
        "level_3_channel_to_kpi_parent": level3,
        "level_4_channel_to_platform": level4,
        "level_5_platform_to_total": total,
        "walmart_badger_gas_control": extraction[
            "walmart_badger_gas_control"
        ],
        "scope_exclusion_control": {
            "expected": scope_exclusion_expected,
            "actual": extraction["excluded_bp_sales"],
            "status": (
                "PASS"
                if scope_exclusion_expected is not None and abs(extraction["excluded_bp_sales"] - scope_exclusion_expected) < 0.005
                else "SKIP_NO_BASELINE"
                if scope_exclusion_expected is None
                else "FAIL"
            ),
            "record_count": len(extraction["excluded_source_rows"]),
            "rule": extraction["scope_definition"]["rule"],
        },
    }

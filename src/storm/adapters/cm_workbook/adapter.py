"""Read-only adapter for the verified 2026 US e-commerce CM workbook.

The adapter deliberately consumes cached formula results (``data_only=True``).
It never recalculates the workbook or rebuilds a competing finance model.
"""

from __future__ import annotations

import hashlib
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

import yaml
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from storm.structured_metrics.schema import (
    BpTargetMonthly,
    CmSnapshot,
    DataException,
    DimSku,
    ExceptionSeverity,
    ImportPlan,
    PeriodType,
    Platform,
    ReconciliationResult,
)


class WorkbookContractError(ValueError):
    """The workbook does not satisfy the inspected source contract."""


MONTHS = {name: number for number, name in enumerate(
    ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"), 1
)}


class CmWorkbookAdapter:
    def __init__(self, config_path: str | Path):
        self.config_path = Path(config_path)
        with self.config_path.open("r", encoding="utf-8") as handle:
            self.config: dict[str, Any] = yaml.safe_load(handle)
        self._validate_config()

    def _validate_config(self) -> None:
        if self.config.get("schema_version") != 1:
            raise WorkbookContractError("unsupported structured metrics config schema_version")
        configured = set(self.config["platforms"])
        expected = {item.value for item in Platform}
        if configured != expected:
            raise WorkbookContractError(f"platform enum mismatch: configured={sorted(configured)}")
        detail = self.config["workbook"]["cm_overview"]["detail_columns"]
        excluded = set(self.config["workbook"]["cm_overview"]["excluded_dfc_columns"])
        if any(excluded.intersection(columns) for columns in detail.values()):
            raise WorkbookContractError("DFC columns cannot appear in a primary platform column set")

    def map_platform(self, raw_value: str) -> Platform:
        for platform_name, definition in self.config["platforms"].items():
            if raw_value in definition["source_values"]:
                return Platform(platform_name)
        raise WorkbookContractError(f"invalid platform mapping: {raw_value!r}")

    def parse(
        self,
        workbook_path: str | Path,
        *,
        snapshot_date: str,
        data_through_date: str,
    ) -> ImportPlan:
        path = Path(workbook_path)
        if not path.is_file():
            raise FileNotFoundError(path)
        file_hash = self._sha256(path)

        values = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            self._validate_workbook_structure(values, formulas)
            derived_bp_classification = self._classify_walmart_bp_overview_formula(formulas)
            mapping, mapping_exceptions, mapping_rows = self._read_sku_mapping(values)
            bp_targets, dim_skus, bp_exceptions, bp_rows, source_year = self._read_bp(values, mapping)
            period_type, selected_month = self._read_period(values)
            observed_data_through = self._derive_data_through(values, source_year, selected_month)
            if observed_data_through != data_through_date:
                raise WorkbookContractError(
                    f"data_through_date mismatch: supplied={data_through_date}, workbook={observed_data_through}"
                )
            cm_snapshots, cm_exceptions, reconciliations, source_totals = self._read_cm(
                values,
                snapshot_date=snapshot_date,
                data_through_date=data_through_date,
                period_type=period_type,
            )
            self._reconcile_bp_facts(
                bp_targets, dim_skus, cm_snapshots, selected_month, period_type,
                reconciliations, source_totals, cm_exceptions, derived_bp_classification,
            )
            source_totals["coverage"] = {
                "MTD": "ACCEPTED" if period_type is PeriodType.MTD else "SOURCE_NOT_AVAILABLE",
                "YTD": "ACCEPTED" if period_type is PeriodType.YTD else "SOURCE_NOT_AVAILABLE",
            }
        finally:
            values.close()
            formulas.close()

        source_period = (
            f"{source_year}-{selected_month:02d}|{period_type.value}"
            if period_type is PeriodType.MTD
            else f"{source_year}-YTD-{selected_month:02d}|{period_type.value}"
        )
        return ImportPlan(
            source_filename=path.name,
            file_hash=file_hash,
            report_type="OFFICIAL_CM_WORKBOOK",
            source_period=source_period,
            snapshot_date=snapshot_date,
            data_through_date=data_through_date,
            input_row_count=mapping_rows + bp_rows + len(cm_snapshots),
            dim_skus=tuple(sorted(dim_skus.values(), key=lambda row: (row.platform.value, row.canonical_sku))),
            bp_targets=tuple(sorted(bp_targets, key=lambda row: (row.year, row.month, row.platform.value, row.canonical_sku))),
            cm_snapshots=tuple(sorted(cm_snapshots, key=lambda row: (row.platform.value, row.brand, row.power_source))),
            exceptions=tuple((*mapping_exceptions, *bp_exceptions, *cm_exceptions)),
            reconciliations=tuple(reconciliations),
            source_totals=source_totals,
        )

    def _validate_workbook_structure(self, values: Any, formulas: Any) -> None:
        sheet_cfg = self.config["workbook"]["sheets"]
        missing = [name for name in sheet_cfg.values() if name not in values.sheetnames]
        if missing:
            raise WorkbookContractError(f"missing required sheets: {missing}")

        mapping_cfg = self.config["workbook"]["sku_mapping"]
        mapping_sheet = values[sheet_cfg["sku_mapping"]]
        actual_headers = [mapping_sheet.cell(mapping_cfg["header_row"], col).value for col in range(1, 6)]
        if actual_headers != mapping_cfg["headers"]:
            raise WorkbookContractError(f"SKU Mapping headers changed: {actual_headers!r}")

        bp_sheet = values[sheet_cfg["bp_targets"]]
        bp_headers = [bp_sheet.cell(1, col).value for col in range(1, 37)]
        required = {
            "SKU", "Customer", "Year", "Month-INT", "Month", "Qty", "TTL amount", "DDP/ALL", "Fixed cost/ALL",
            "MKT-Insite/ALL", "MKT-Offsite(种草)/ALL", "MKT-Offsite(Channel MKT)/ALL",
            "Return+Warranty/ALL", "Funding/ALL",
        }
        if not required.issubset(set(bp_headers)):
            raise WorkbookContractError(f"KPI Rawdata missing required headers: {sorted(required - set(bp_headers))}")

        overview_values = values[sheet_cfg["cm_overview"]]
        overview_formulas = formulas[sheet_cfg["cm_overview"]]
        if overview_formulas["I82"].value != "=G82+H82":
            raise WorkbookContractError("Overview THD formula no longer exposes Main + DFC at I82")
        if overview_values["L258"].value != "The Home Depot Inc DFC" or overview_values["M258"].value != "The Home Depot Inc DFC":
            raise WorkbookContractError("Overview DFC detail columns changed")

    def _classify_walmart_bp_overview_formula(self, formulas: Any) -> str:
        """Classify the inspected Walmart Sunseeker BP derived-view discrepancy."""
        sheet = formulas[self.config["workbook"]["sheets"]["cm_overview"]]
        expected_customer = next(
            raw for raw, platform in self.config["workbook"]["bp_targets"]["scoped_customers"].items()
            if platform == Platform.WALMART_MP.value
        )
        formula = self._text(sheet["P300"].value)
        if (
            sheet["P296"].value == "=P258"
            and self._text(sheet["P258"].value) != expected_customer
            and "$D:$D,P$296" in formula
            and self._text(sheet["Z258"].value) == expected_customer
        ):
            return "BROKEN_DERIVED_FORMULA"
        if formula.startswith("="):
            return "UNRESOLVED"
        return "UNRESOLVED"

    def _read_sku_mapping(self, workbook: Any) -> tuple[dict[str, dict[str, Any]], list[DataException], int]:
        sheet_name = self.config["workbook"]["sheets"]["sku_mapping"]
        sheet = workbook[sheet_name]
        mapping: dict[str, dict[str, Any]] = {}
        exceptions: list[DataException] = []
        rows = 0
        for row_number, row in enumerate(sheet.iter_rows(min_row=2, max_col=5, values_only=True), 2):
            if not any(value is not None for value in row):
                continue
            rows += 1
            sku, asin, power_source, brand, category = row
            sku = self._text(sku)
            if not sku:
                exceptions.append(DataException(
                    "INVALID_SKU_MAPPING", "mapping row has no SKU", source_sheet=sheet_name, source_row=row_number
                ))
                continue
            record = {
                "asin": self._text(asin) or None,
                "power_source": self._text(power_source),
                "brand": self._text(brand),
                "category": self._text(category) or None,
                "source_row": row_number,
            }
            if not record["power_source"] or not record["brand"]:
                exceptions.append(DataException(
                    "INVALID_SKU_MAPPING", "mapping row lacks Brand or Pource Source",
                    raw_key=sku, source_sheet=sheet_name, source_row=row_number,
                ))
                continue
            if sku in mapping:
                comparable = {key: value for key, value in record.items() if key != "source_row"}
                existing = {key: value for key, value in mapping[sku].items() if key != "source_row"}
                conflicting = existing != comparable
                exceptions.append(DataException(
                    "DUPLICATE_SKU_MAPPING",
                    "conflicting duplicate SKU mapping" if conflicting else "duplicate normalized SKU mapping",
                    severity=ExceptionSeverity.ERROR if conflicting else ExceptionSeverity.WARNING,
                    raw_key=sku, source_sheet=sheet_name, source_row=row_number,
                ))
                continue
            mapping[sku] = record
        return mapping, exceptions, rows

    def _read_bp(
        self, workbook: Any, mapping: dict[str, dict[str, Any]]
    ) -> tuple[list[BpTargetMonthly], dict[tuple[Platform, str], DimSku], list[DataException], int, int]:
        sheet_name = self.config["workbook"]["sheets"]["bp_targets"]
        sheet = workbook[sheet_name]
        headers = [self._text(cell.value) for cell in sheet[1]]
        index = {name: position for position, name in enumerate(headers)}
        scoped = self.config["workbook"]["bp_targets"]["scoped_customers"]
        aggregate: dict[tuple[int, int, Platform, str], dict[str, Any]] = {}
        dim_skus: dict[tuple[Platform, str], DimSku] = {}
        exceptions: list[DataException] = []
        source_years: set[int] = set()
        input_rows = 0

        components = ["TTL amount", "DDP/ALL", "Fixed cost/ALL", "MKT-Insite/ALL", "MKT-Offsite(种草)/ALL",
                      "MKT-Offsite(Channel MKT)/ALL", "Return+Warranty/ALL", "Funding/ALL"]
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if not any(value is not None for value in values):
                continue
            input_rows += 1
            customer = self._text(values[index["Customer"]])
            if customer not in scoped:
                continue
            try:
                platform = Platform(scoped[customer])
            except ValueError:
                exceptions.append(DataException(
                    "INVALID_PLATFORM", f"unsupported configured platform {scoped[customer]!r}",
                    raw_key=customer, source_sheet=sheet_name, source_row=row_number,
                ))
                continue
            sku = self._text(values[index["SKU"]])
            mapped = mapping.get(sku)
            if mapped is None:
                exceptions.append(DataException(
                    "UNMAPPED_SKU", "target BP row has no verified SKU Mapping record",
                    platform=platform.value, raw_key=sku, source_sheet=sheet_name, source_row=row_number,
                ))
                continue
            try:
                year = int(values[index["Year"]])
                month_header = self.config["workbook"]["bp_targets"]["month_number_header"]
                month = int(values[index[month_header]])
            except (TypeError, ValueError):
                exceptions.append(DataException(
                    "INVALID_PERIOD", "target BP row has invalid Year or Month",
                    platform=platform.value, raw_key=sku, source_sheet=sheet_name, source_row=row_number,
                ))
                continue
            if not 1 <= month <= 12:
                exceptions.append(DataException(
                    "INVALID_PERIOD", f"Month outside 1..12: {month}", platform=platform.value,
                    raw_key=sku, source_sheet=sheet_name, source_row=row_number,
                ))
                continue
            source_years.add(year)
            key = (year, month, platform, sku)
            bucket = aggregate.setdefault(key, {"units": [], "sales": [], "cm": [], "rows": []})
            units = self._number_or_none(values[index["Qty"]])
            component_values = [self._number_or_none(values[index[name]]) for name in components]
            sales = component_values[0]
            cm = None if any(item is None for item in component_values) else (
                component_values[0] - sum(component_values[1:])
            )
            bucket["units"].append(units)
            bucket["sales"].append(sales)
            bucket["cm"].append(cm)
            bucket["rows"].append(row_number)
            platform_cfg = self.config["platforms"][platform.value]
            dim_skus[(platform, sku)] = DimSku(
                platform=platform,
                canonical_sku=sku,
                platform_sku=None,
                brand=mapped["brand"],
                power_source=mapped["power_source"],
                category=mapped["category"],
                business_model=platform_cfg["business_model"],
                primary_sales_basis=platform_cfg["primary_sales_basis"],
                source_row=mapped["source_row"],
            )

        if len(source_years) != 1:
            raise WorkbookContractError(f"target BP data must expose exactly one year, got {sorted(source_years)}")
        bp_targets: list[BpTargetMonthly] = []
        for (year, month, platform, sku), bucket in aggregate.items():
            units = self._sum_nullable(bucket["units"])
            sales = self._sum_nullable(bucket["sales"])
            cm = self._sum_nullable(bucket["cm"])
            bp_targets.append(BpTargetMonthly(
                year=year,
                month=month,
                platform=platform,
                canonical_sku=sku,
                bp_units=units,
                bp_sales=sales,
                bp_cm=cm,
                bp_cm_pct=None if sales in (None, 0) or cm is None else cm / sales,
                source_sheet=sheet_name,
                source_rows=tuple(bucket["rows"]),
            ))
        return bp_targets, dim_skus, exceptions, input_rows, next(iter(source_years))

    def _read_period(self, workbook: Any) -> tuple[PeriodType, int]:
        cfg = self.config["workbook"]["cm_overview"]
        sheet = workbook[self.config["workbook"]["sheets"]["cm_overview"]]
        month_name = self._text(sheet[cfg["selector_month_cell"]].value)
        period_name = self._text(sheet[cfg["selector_period_cell"]].value).upper()
        if month_name not in MONTHS:
            raise WorkbookContractError(f"unsupported Overview month selector: {month_name!r}")
        try:
            period_type = PeriodType(period_name)
        except ValueError as exc:
            raise WorkbookContractError(f"unsupported Overview period selector: {period_name!r}") from exc
        return period_type, MONTHS[month_name]

    def _derive_data_through(self, workbook: Any, year: int, month: int) -> str:
        if "Actual Orders" not in workbook.sheetnames:
            raise WorkbookContractError("Actual Orders sheet missing")
        sheet = workbook["Actual Orders"]
        headers = [self._text(cell.value) for cell in sheet[1]]
        index = {name: position for position, name in enumerate(headers)}
        needed = {"Filter", "Year", "Month", "Date", "Channel"}
        if not needed.issubset(index):
            raise WorkbookContractError(f"Actual Orders missing headers: {sorted(needed - set(index))}")
        scoped = set(self.config["workbook"]["bp_targets"]["scoped_customers"])
        dates: list[date] = []
        for values in sheet.iter_rows(min_row=2, values_only=True):
            if values[index["Filter"]] != 1:
                continue
            if values[index["Year"]] != year or values[index["Month"]] != month:
                continue
            if self._text(values[index["Channel"]]) not in scoped:
                continue
            raw_date = values[index["Date"]]
            if isinstance(raw_date, datetime):
                dates.append(raw_date.date())
            elif isinstance(raw_date, date):
                dates.append(raw_date)
        if not dates:
            raise WorkbookContractError("no scoped Actual Orders rows for selected period")
        return max(dates).isoformat()

    def _read_cm(
        self,
        workbook: Any,
        *,
        snapshot_date: str,
        data_through_date: str,
        period_type: PeriodType,
    ) -> tuple[list[CmSnapshot], list[DataException], list[ReconciliationResult], dict[str, Any]]:
        sheet_name = self.config["workbook"]["sheets"]["cm_overview"]
        sheet = workbook[sheet_name]
        cfg = self.config["workbook"]["cm_overview"]
        actual_rows = cfg["actual_rows"]
        bp_rows = cfg["bp_rows"]
        dim_rows = cfg["dimension_rows"]
        snapshots: list[CmSnapshot] = []
        exceptions: list[DataException] = []
        reconciliations: list[ReconciliationResult] = []

        for platform_name, columns in cfg["detail_columns"].items():
            platform = Platform(platform_name)
            platform_cfg = self.config["platforms"][platform_name]
            for column in columns:
                brand = self._text(sheet[f"{column}{dim_rows['brand']}"].value)
                power_source = self._text(sheet[f"{column}{dim_rows['power_source']}"].value)
                source_platform = self._text(sheet[f"{column}{dim_rows['platform']}"].value)
                if not brand or not power_source:
                    exceptions.append(DataException(
                        "INVALID_CM_DIMENSION", "Overview detail column lacks Brand or Power Source",
                        platform=platform.value, source_sheet=sheet_name,
                        source_range=f"{column}{dim_rows['platform']}:{column}{bp_rows['cm_pct']}",
                    ))
                    continue
                if platform is Platform.THD and "DFC" in source_platform.upper():
                    raise WorkbookContractError(f"DFC column leaked into THD primary set: {column}")

                actual = {name: self._metric(sheet[f"{column}{row}"].value, exceptions, platform, sheet_name, f"{column}{row}")
                          for name, row in actual_rows.items()}
                bp = {name: self._metric(sheet[f"{column}{row}"].value, exceptions, platform, sheet_name, f"{column}{row}")
                      for name, row in bp_rows.items()}
                snapshot = CmSnapshot(
                    snapshot_date=snapshot_date,
                    data_through_date=data_through_date,
                    period_type=period_type,
                    platform=platform,
                    brand=brand,
                    power_source=power_source,
                    actual_units=actual["units"], actual_sales=actual["sales"], actual_gm=actual["gm"],
                    actual_gm_pct=self._ratio(actual["gm"], actual["sales"]), actual_cm=actual["cm"],
                    actual_cm_pct=self._ratio(actual["cm"], actual["sales"]),
                    bp_units=bp["units"], bp_sales=bp["sales"], bp_cm=bp["cm"],
                    bp_cm_pct=self._ratio(bp["cm"], bp["sales"]),
                    cm_basis="PRIMARY_SELL_IN_EXCLUDES_DFC" if platform is Platform.THD else "PRIMARY",
                    primary_sales_basis=platform_cfg["primary_sales_basis"],
                    source_sheet=sheet_name,
                    source_range=f"{column}{dim_rows['platform']}:{column}{bp_rows['cm_pct']}",
                )
                snapshots.append(snapshot)
                for metric in ("actual_units", "actual_sales", "actual_cm", "bp_units", "bp_sales", "bp_cm"):
                    value = getattr(snapshot, metric)
                    reconciliations.append(self._reconcile(
                        "BRAND_POWER", f"{platform.value}|{brand}|{power_source}", metric,
                        value, value, self._tolerance(metric),
                    ))

        platform_totals: dict[str, Any] = {}
        for platform in Platform:
            selected = [row for row in snapshots if row.platform is platform]
            platform_totals[platform.value] = {}
            summary_columns = cfg["platform_summary"]["columns"][platform.value]
            for kind in ("actual", "bp"):
                for metric in ("units", "sales", "cm"):
                    attribute = f"{kind}_{metric}"
                    imported = self._sum_nullable([getattr(row, attribute) for row in selected])
                    summary_row = cfg["platform_summary"][kind][metric]
                    official = self._sum_nullable([
                        self._number_or_none(sheet[f"{column}{summary_row}"].value) for column in summary_columns
                    ])
                    result = self._reconcile(
                        "PLATFORM", platform.value, attribute, official, imported, self._tolerance(attribute)
                    )
                    reconciliations.append(result)
                    platform_totals[platform.value][attribute] = {
                        "official": official, "imported": imported, "variance": result.variance,
                        "tolerance": result.tolerance, "result": result.result,
                    }
        return snapshots, exceptions, reconciliations, {"platform": platform_totals}

    def _reconcile_bp_facts(
        self,
        bp_targets: list[BpTargetMonthly],
        dim_skus: dict[tuple[Platform, str], DimSku],
        cm_snapshots: list[CmSnapshot],
        selected_month: int,
        period_type: PeriodType,
        reconciliations: list[ReconciliationResult],
        source_totals: dict[str, Any],
        exceptions: list[DataException],
        derived_view_classification: str,
    ) -> None:
        """Cross-check imported KPI Rawdata facts against official Overview BP caches."""
        grouped: dict[tuple[Platform, str, str], dict[str, list[float | None]]] = defaultdict(
            lambda: {"bp_units": [], "bp_sales": [], "bp_cm": []}
        )
        for row in bp_targets:
            in_period = row.month == selected_month if period_type is PeriodType.MTD else row.month <= selected_month
            if not in_period:
                continue
            dimension = dim_skus[(row.platform, row.canonical_sku)]
            bucket = grouped[(row.platform, dimension.brand, dimension.power_source)]
            bucket["bp_units"].append(row.bp_units)
            bucket["bp_sales"].append(row.bp_sales)
            bucket["bp_cm"].append(row.bp_cm)

        detail_totals: dict[str, Any] = {}
        snapshot_keys = {(row.platform, row.brand, row.power_source): row for row in cm_snapshots}
        for key, snapshot in snapshot_keys.items():
            platform, brand, power_source = key
            bucket = grouped.get(key, {"bp_units": [], "bp_sales": [], "bp_cm": []})
            result_key = f"{platform.value}|{brand}|{power_source}"
            detail_totals[result_key] = {}
            for metric in ("bp_units", "bp_sales", "bp_cm"):
                imported = self._sum_nullable(bucket[metric]) if bucket[metric] else 0.0
                official = getattr(snapshot, metric)
                result = self._reconcile(
                    "BP_DERIVED_VIEW_BRAND_POWER", result_key, metric, official, imported,
                    self._tolerance(metric), is_blocking=False,
                    classification=derived_view_classification if official != imported else None,
                )
                reconciliations.append(result)
                if result.result != "PASS":
                    exceptions.append(DataException(
                        derived_view_classification,
                        f"authoritative KPI Rawdata {metric}={imported} differs from secondary Overview {metric}={official}",
                        severity=ExceptionSeverity.WARNING,
                        platform=platform.value,
                        raw_key=f"{brand}|{power_source}|{metric}",
                        source_sheet="KPI Rawdata / Overview",
                    ))
                detail_totals[result_key][metric] = {
                    "official": official, "imported": imported, "variance": result.variance,
                    "tolerance": result.tolerance, "result": result.result,
                }

        platform_totals: dict[str, Any] = {}
        for platform in Platform:
            platform_totals[platform.value] = {}
            relevant = [row for row in cm_snapshots if row.platform is platform]
            for metric in ("bp_units", "bp_sales", "bp_cm"):
                official = self._sum_nullable([getattr(row, metric) for row in relevant])
                imported_values = [
                    self._sum_nullable(bucket[metric])
                    for (item_platform, _, _), bucket in grouped.items()
                    if item_platform is platform
                ]
                imported = self._sum_nullable(imported_values) if imported_values else 0.0
                result = self._reconcile(
                    "BP_DERIVED_VIEW_PLATFORM", platform.value, metric, official, imported,
                    self._tolerance(metric), is_blocking=False,
                    classification=derived_view_classification if official != imported else None,
                )
                reconciliations.append(result)
                if result.result != "PASS":
                    exceptions.append(DataException(
                        derived_view_classification,
                        f"authoritative KPI Rawdata platform {metric}={imported} differs from secondary Overview {metric}={official}",
                        severity=ExceptionSeverity.WARNING,
                        platform=platform.value,
                        raw_key=f"PLATFORM|{metric}",
                        source_sheet="KPI Rawdata / Overview",
                    ))
                platform_totals[platform.value][metric] = {
                    "official": official, "imported": imported, "variance": result.variance,
                    "tolerance": result.tolerance, "result": result.result,
                }
        source_totals["bp_fact_brand_power"] = detail_totals
        source_totals["bp_fact_platform"] = platform_totals

    def _metric(
        self, value: Any, exceptions: list[DataException], platform: Platform, sheet_name: str, cell: str
    ) -> float | None:
        try:
            return self._number_or_none(value)
        except WorkbookContractError:
            exceptions.append(DataException(
                "INVALID_CACHED_VALUE", f"formula cache is not numeric: {value!r}",
                platform=platform.value, source_sheet=sheet_name, source_range=cell,
            ))
            return None

    def _tolerance(self, metric: str) -> float:
        cfg = self.config["reconciliation"]
        if "units" in metric:
            return float(cfg["unit_tolerance"])
        if "pct" in metric:
            return float(cfg["ratio_tolerance"])
        return float(cfg["currency_tolerance"])

    @staticmethod
    def _reconcile(
        level: str,
        key: str,
        metric: str,
        official: float | None,
        imported: float | None,
        tolerance: float,
        *,
        is_blocking: bool = True,
        classification: str | None = None,
    ) -> ReconciliationResult:
        if official is None and imported is None:
            variance, result = None, "PASS"
        elif official is None or imported is None:
            variance, result = None, "FAIL"
        else:
            variance = imported - official
            result = "PASS" if abs(variance) <= tolerance else "FAIL"
        return ReconciliationResult(
            level, key, metric, official, imported, variance, tolerance, result,
            is_blocking, classification,
        )

    @staticmethod
    def _ratio(numerator: float | None, denominator: float | None) -> float | None:
        return None if numerator is None or denominator in (None, 0) else numerator / denominator

    @staticmethod
    def _sum_nullable(values: Iterable[float | None]) -> float | None:
        materialized = list(values)
        return None if any(value is None for value in materialized) else sum(materialized)

    @staticmethod
    def _number_or_none(value: Any) -> float | None:
        if value is None or value == "":
            return None
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise WorkbookContractError(f"expected numeric cached value, got {value!r}")
        return float(value)

    @staticmethod
    def _text(value: Any) -> str:
        return "" if value is None else str(value).strip()

    @staticmethod
    def _sha256(path: Path) -> str:
        digest = hashlib.sha256()
        with path.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest().upper()
